"""
exel_export.py - Экспорт ведомости работ в Excel
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from datetime import datetime
from generator import GeneratorВедомости
from works_data import РаботыБаза

class ExcelExport:
    def __init__(self, настройки):
        """
        настройки: словарь с параметрами из GUI
        цвета: цвет_шапки, цвет_основной, цвет_второстепенный
        путь_сохранения, формат_имени
        """
        self.настройки = настройки
        self.генератор = GeneratorВедомости()
        self.работы = РаботыБаза().работы_база

    def сохранить_в_excel(self, выбранные_работы, путь=None):
        """
        выбранные_работы: список ключей из базы работ (str)
        путь: путь сохранения файла (по умолчанию из настроек)
        """
        # Создаем книгу и лист
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ведомость работ"

        # Цвета
        fill_шапка = PatternFill(start_color=self.настройки['цвет_шапки'], end_color=self.настройки['цвет_шапки'], fill_type="solid")
        fill_основной = PatternFill(start_color=self.настройки['цвет_основной'], end_color=self.настройки['цвет_основной'], fill_type="solid")
        fill_второстепенный = PatternFill(start_color=self.настройки['цвет_второстепенный'], end_color=self.настройки['цвет_второстепенный'], fill_type="solid")
        
        # Заголовки
        заголовки = ["№", "Раздел", "Вид работы", "Формула", "Ед.изм.", "Значение"]
        ws.append(заголовки)
        for col in range(1, len(заголовки)+1):
            ws.cell(row=1, column=col).fill = fill_шапка
            ws.cell(row=1, column=col).font = Font(bold=True)
            ws.cell(row=1, column=col).alignment = Alignment(horizontal='center', vertical='center')

        текущий_порядок = 1
        row_idx = 2

        for работа_key in выбранные_работы:
            if работа_key not in self.работы:
                continue
            работа = self.работы[работа_key]
            раздел = работа["раздел"]
            for item in работа["работы"]:
                вид = item["вид"]

                # Подставляем размеры плитки, если есть
                if "размер" in item and item["размер"]:
                    размер = self.настройки.get("размер_плитки", "60x60")
                    вид = вид.format(размер=размер)

                формула = item.get("формула", "")
                ед = item.get("ед", "")
                значение = self.генератор.рассчитать(формула)

                ws.append([текущий_порядок, раздел, вид, формула, ед, значение])

                # Цвет строк: основной/второстепенный чередование
                fill = fill_основной if текущий_порядок % 2 else fill_второстепенный
                for col in range(1, 7):
                    ws.cell(row=row_idx, column=col).fill = fill
                row_idx += 1
                текущий_порядок += 1

        # Автоподбор ширины колонок
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max_length + 2

        # Имя файла
        if путь is None:
            путь = self.настройки.get("путь_сохранения", "")
        имя_файла = datetime.now().strftime(self.настройки.get("формат_имени", "Ведомость_%Y%m%d_%H%M")) + ".xlsx"
        полный_путь = путь + "/" + имя_файла if путь else имя_файла

        wb.save(полный_путь)
        return полный_путь