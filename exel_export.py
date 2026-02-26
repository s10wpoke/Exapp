from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def сохранить_в_excel(строки, имя_файла):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ведомость"

    # ==== ТВОЙ КОД ФОРМАТИРОВАНИЯ ====
    # границы
    # цвета
    # ширина колонок
    # объединения
    # стили
    # заполнение строк

    for row_index, строка in enumerate(строки, start=1):
        ws.cell(row=row_index, column=1, value=строка[0])
        ws.cell(row=row_index, column=2, value=строка[1])
        ws.cell(row=row_index, column=3, value=строка[2])
        ws.cell(row=row_index, column=4, value=строка[3])

    wb.save(имя_файла)